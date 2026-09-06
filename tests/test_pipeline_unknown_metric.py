"""Сквозной прогон ноды на реальном run_package с подменённым LLM.

Три сценария «новый агент»:
1. отчёт называет macro-F1 — модель записывает формулу как в отчёте, пересчёт
   на корзине совпадает → computed, формула уходит в контракт;
2. модель подобрала не ту формулу (среднее колонки вместо F1) — пересчёт не
   совпадает с отчётом → not_computable с кодом km_reconciliation_mismatch;
3. контроль: готовый метод, воспроизводящий число отчёта → computed.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

import main as node
from laim_basket.pipeline import run_package

BASKET_ID = "CI00000001"


def _write_docx(path: Path, paragraphs: list[str]) -> None:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(path)


def _write_package(root: Path, reported: str, scores: list[int]) -> Path:
    package = root / BASKET_ID
    package.mkdir()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Лист1"
    sheet.append(["query_id", "input_query", "Класс агента", "Истинный класс", "Итог"])
    classes = [("a", "a"), ("a", "b"), ("b", "b"), ("b", "b"), ("a", "b")]
    for index, score in enumerate(scores, start=1):
        agent, truth = classes[(index - 1) % len(classes)]
        sheet.append([f"q{index}", f"вопрос {index}", agent, truth, score])
    workbook.save(package / "test_set.xlsx")
    _write_docx(package / "assessor_instruction.docx", [
        "Инструкция ассессора: каждая пара запрос-ответ оценивается независимо.",
        "Итог равен 1, если ответ верный, иначе 0.",
    ])
    _write_docx(package / "development_report.docx", [
        "Отчёт о разработке агента.",
        "Ключевая метрика качества описана в отчёте о валидации.",
    ])
    _write_docx(package / "validation_report.docx", [
        "Отчёт о валидации агента.",
        "Ключевая метрика: macro-F1 по классам на тестовой корзине.",
        f"Значение метрики: {reported}",
    ])
    return package


LAYOUT = {
    "layout_version": "laim-layout.v1",
    "basket_id": BASKET_ID,
    "sheet_name": "Лист1",
    "ignored_sheets": [],
    "header_rows": [1],
    "roles": {
        "input_query": {"column": "B", "header": "input_query"},
        "output_answer": {"column": "C", "header": "Класс агента"},
        "query_id": {"column": "A", "header": "query_id"},
        "scenario": None,
        "assessor_id": None,
        "reference_answers": [],
    },
    "grouping": {"kind": "none", "column": None},
    "dialogue_blob": None,
    "weight": None,
}


def _measurement(reported: str, formula: str | None = None) -> dict:
    if formula:
        return {
            **_measurement(reported),
            "metric_name": "Macro F1",
            "score": {
                "method": "formula",
                "sources": [
                    {"column_id": "C", "name": "prediction", "role": "prediction",
                     "normalization": "label", "polarity": "direct"},
                    {"column_id": "D", "name": "target", "role": "target",
                     "normalization": "label", "polarity": "direct"},
                ],
                "missing_policy": "fail",
                "majority_denominator": None,
            },
            "formula": formula,
            "release": {"threshold": None, "comparator": None, "scale": "ratio", "precision": 4},
        }
    return {
        "plan_version": "laim-measurement-plan.v2",
        "basket_id": BASKET_ID,
        "metric_name": "F1",
        "document_roles": {
            "instruction": "doc-1",
            "development_report": "doc-2",
            "validation_report": "doc-3",
        },
        "assessment_mode": "qa",
        "score": {
            "method": "identity",
            "sources": [{
                "column_id": "E", "role": "final_score",
                "normalization": "numeric", "polarity": "direct",
            }],
            "missing_policy": "fail",
            "majority_denominator": None,
        },
        "formula": None,
        "reducer": {"method": "mean"},
        "release": {"threshold": None, "comparator": None, "scale": "ratio", "precision": 2},
        "reported_value_state": "unambiguous",
        "reported_value": {"value": reported, "raw": reported, "span_id": "doc-3:p0003"},
        "evidence": {
            "metric": ["doc-3:p0002"],
            "score": ["doc-1:p0002"],
            "assessment_mode": ["doc-1:p0001"],
            "missing_policy": [],
            "denominator": [],
            "reducer": ["doc-3:p0002"],
            "release": [],
            "reported_value": ["doc-3:p0003"],
        },
    }


class ScriptedLlm:
    """Возвращает один и тот же layout и план: модели больше нечего предложить."""

    def __init__(self, reported: str, formula: str | None = None):
        self.reported = reported
        self.formula = formula
        self.calls: list[str] = []

    def chat(self, messages: list[dict], label: str) -> str:
        self.calls.append(label)
        proposal = LAYOUT if label.startswith("layout") else _measurement(self.reported, self.formula)
        return json.dumps(proposal, ensure_ascii=False)


def _run(tmp_path: Path, reported: str, scores: list[int], formula: str | None = None):
    package = _write_package(tmp_path, reported, scores)
    client = ScriptedLlm(reported, formula)
    result = run_package(package, tmp_path / "out", client=client)
    return result, client


def test_report_metric_is_written_as_formula_and_reproduced(tmp_path):
    # классы (агент, истина) × 5: a/a, a/b, b/b, b/b, a/b → macro-F1 = 7/12 ≈ 0.5833
    result, client = _run(
        tmp_path, reported="0.5833", scores=[1, 0, 1, 1, 0], formula='f1(prediction, target, "macro")',
    )
    assert result.status == "computed"
    assert result.km["reconciliation"]["status"] == "match"
    assert sum(label.startswith("measurement") for label in client.calls) == 1

    contract = node._monitoring_metric(result)
    assert contract["status"] == "computed"
    assert contract["formula"] == 'f1(prediction, target, "macro")'
    assert contract["scoring"]["method"] == "formula"
    assert [s["name"] for s in contract["scoring"]["sources"]] == ["prediction", "target"]
    assert contract["baseline"]["value"] == pytest.approx(0.5833)
    assert contract["baseline"]["recomputed_value"] == pytest.approx(7 / 12)


def test_wrong_formula_does_not_reproduce_report_and_is_refused(tmp_path):
    # модель взяла среднее колонки Итог (0.6) вместо macro-F1 из отчёта (0.5833)
    result, client = _run(tmp_path, reported="0.5833", scores=[1, 0, 1, 1, 0])
    assert result.status == "not_evaluable"
    assert result.km["reason_code"] == "km_reconciliation_mismatch"
    assert result.km["details"]["recomputed_value"] == pytest.approx(0.6)
    assert sum(label.startswith("measurement") for label in client.calls) == 5

    contract = node._monitoring_metric(result)
    assert contract["status"] == "not_computable"
    assert contract["reason_code"] == "km_reconciliation_mismatch"


def test_preset_method_reproducing_report_is_published(tmp_path):
    result, _client = _run(tmp_path, reported="0.75", scores=[1, 1, 1, 0])
    assert result.status == "computed"
    contract = node._monitoring_metric(result)
    assert contract["status"] == "computed"
    assert contract["formula"] == "mean(source_1)"
    assert contract["baseline"]["value"] == pytest.approx(0.75)
    assert contract["baseline"]["reconciliation"] == "match"
    assert contract["scoring"]["sources"][0]["column_name"] == "итог_metric"
